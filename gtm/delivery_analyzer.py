"""
Delivery Document Analyzer
--------------------------
Extracts text from uploaded business documents (CSV, XLSX, PDF, DOCX, images,
TXT, JSON) and uses a three-stage NLP/ML pipeline + Gemini to auto-score the
6 Delivery assessment questions.

Pipeline
--------
Stage 1 — Text extraction   : file-format parsers (csv, openpyxl, pypdf, etc.)
Stage 2 — NLP / ML layer    :
    A. TF-IDF (scikit-learn)       — keyword-ranked paragraphs per question
    B. Sentence-transformers       — semantically similar sentences per question
    C. spaCy NER                   — quantitative metrics & entities extracted
Stage 3 — LLM scoring       : enriched, structured evidence fed to Gemini-2.5-flash

Human-AI design principle: the agent proposes scores grounded in document
evidence; the human reviews and overrides before submitting the assessment.
"""

import csv
import json
import logging
import re
import threading
import zipfile
import xml.etree.ElementTree as ET

logger = logging.getLogger(__name__)

# Per-document text limit (chars) sent to Gemini.
_TEXT_LIMIT_PER_DOC = 30_000
# Combined text limit across all documents.
_COMBINED_TEXT_LIMIT = 50_000

# Lazy-loaded ML model caches — initialised on first analysis call, then reused.
# Locks guard against two threads in the same worker racing to load the same
# model concurrently (mirrors gtm/ai_services.py's _get_client() pattern).
_ST_MODEL = None    # sentence-transformers SentenceTransformer
_ST_LOCK = threading.Lock()
_SPACY_NLP = None   # spaCy en_core_web_sm pipeline (tok2vec + ner only, see _get_spacy_nlp)
_SPACY_LOCK = threading.Lock()


DELIVERY_QUESTIONS = [
    {
        "id_code": "DEL-TTV-01",
        "dimension": "Time to Value",
        "text": (
            "We track how long it takes new customers to get their first real value, "
            "by segment, and we actively work to shorten that time."
        ),
        "evidence_hint": (
            "Look for: time-to-value (TTV) metrics, activation timelines by customer segment, "
            "onboarding completion rates, TTV improvement initiatives, first-value milestone data."
        ),
    },
    {
        "id_code": "DEL-ONB-02",
        "dimension": "Onboarding Discipline",
        "text": (
            "Our onboarding process has clear milestones, clear owners, and realistic "
            "completion targets, and progress is tracked in one shared system."
        ),
        "evidence_hint": (
            "Look for: onboarding checklists, milestone definitions with owners, "
            "completion-date targets, project or task-tracker data, onboarding process documentation."
        ),
    },
    {
        "id_code": "DEL-HLT-03",
        "dimension": "Health Monitoring",
        "text": (
            "We score customer health using product usage, engagement, and support signals, "
            "and we use clear playbooks to act early on at-risk accounts."
        ),
        "evidence_hint": (
            "Look for: health-score models or rubrics, usage/engagement metrics, "
            "support-ticket trend data, at-risk account lists, CS intervention playbooks."
        ),
    },
    {
        "id_code": "DEL-RET-04",
        "dimension": "Retention",
        "text": (
            "We review gross and net retention by cohort every month, and we launch "
            "targeted actions quickly when any segment starts to decline."
        ),
        "evidence_hint": (
            "Look for: gross/net retention rates, cohort tables, monthly review cadence evidence, "
            "churn data by segment, save-motion records or win-back campaigns."
        ),
    },
    {
        "id_code": "DEL-QBR-05",
        "dimension": "Success Governance",
        "text": (
            "Our high-value customers receive regular business reviews focused on outcomes, "
            "roadmap alignment, and practical expansion opportunities."
        ),
        "evidence_hint": (
            "Look for: QBR schedules or meeting cadence, business-review agendas or notes, "
            "customer outcome tracking, expansion pipeline linked to reviews."
        ),
    },
    {
        "id_code": "DEL-ADV-06",
        "dimension": "Advocacy",
        "text": (
            "After customers achieve clear value, we consistently capture proof points like "
            "quotes, case studies, and references to support future selling."
        ),
        "evidence_hint": (
            "Look for: testimonial requests or a testimonial pipeline, case-study drafts, "
            "reference customer lists, an advocacy or customer-marketing programme."
        ),
    },
]


# ---------------------------------------------------------------------------
# Fallback guidance — used only when Gemini is unavailable (see
# _scoring_memory_fallback below). Written once per question x score-tier,
# reviewed for quality, never referencing any specific company's evidence —
# paired at runtime with the *user's own* evidence excerpt, which is safe
# because it's their own document, not another company's.
# ---------------------------------------------------------------------------
QUESTION_SCORE_GUIDANCE = {
    "DEM-ICP-01": {
        "low": "Without a written ICP, marketing and sales tend to chase whichever leads show up rather than the ones most likely to close and stay — that scatter shows up later as low win rates and early churn.",
        "mid": "An ICP exists but isn't consistently reviewed or applied, so teams may drift back to gut-feel targeting between updates.",
        "high": "A clearly documented and regularly reviewed ICP keeps marketing, sales, and product aligned on exactly who you're building for.",
    },
    "DEM-FIT-02": {
        "low": "Low ICP-fit rates on inbound leads usually mean top-of-funnel messaging or targeting is pulling in the wrong audience, which shows up later as poor sales conversion.",
        "mid": "Some leads match the ICP but qualification fields aren't consistently filled in, making it hard to verify fit rates with confidence.",
        "high": "A high, verifiable share of inbound leads matching the ICP is a strong signal that targeting and messaging are working together.",
    },
    "DEM-MSG-03": {
        "low": "Inconsistent or untested messaging across channels usually means prospects hear a different value story depending on where they encounter you, which slows down buying decisions.",
        "mid": "A messaging framework exists but isn't consistently used across every channel or tested with real segments.",
        "high": "Consistent, tested messaging across website, outbound, and sales materials shortens the path from first touch to understood value.",
    },
    "DEM-CHN-04": {
        "low": "Without visibility into channel-level CAC and pipeline contribution, budget tends to get allocated by habit rather than performance.",
        "mid": "Channel performance is tracked in some form but not reviewed on a regular cadence, so underperforming spend can linger.",
        "high": "A monthly-reviewed channel plan with CAC and pipeline data by source lets you double down on what's actually working.",
    },
    "DEM-ATT-05": {
        "low": "Without reliable attribution, it's hard to tell which channels are actually driving pipeline — decisions end up based on the loudest channel, not the most effective one.",
        "mid": "Some attribution data exists but isn't fully trusted or consistently used to inform channel decisions.",
        "high": "Reliable attribution reporting means marketing spend can be defended and reallocated with real data, not guesswork.",
    },
    "DEM-CNT-06": {
        "low": "Ad-hoc content or outbound activity, disconnected from target accounts and pipeline goals, rarely compounds into predictable pipeline.",
        "mid": "A schedule exists but isn't consistently tied to specific target accounts or pipeline goals.",
        "high": "A steady, goal-tied content or outbound cadence is one of the most reliable ways to keep pipeline predictable.",
    },
    "CON-SLA-01": {
        "low": "Without response-time targets, high-intent leads often go cold before a rep ever reaches out — this is one of the most common, most fixable pipeline leaks.",
        "mid": "Response-time targets exist but aren't consistently reviewed, so breaches can go unnoticed.",
        "high": "Clear SLA targets with weekly review keep response times tight, which reliably lifts lead-to-opportunity conversion.",
    },
    "CON-QLF-02": {
        "low": "Without mandatory qualification gates, unqualified opportunities advance through the pipeline and distort forecasts and rep time allocation.",
        "mid": "A qualification framework exists but isn't consistently enforced at every stage gate.",
        "high": "Enforced qualification gates keep the pipeline honest and free up rep time for deals that can actually close.",
    },
    "CON-STG-03": {
        "low": "Without documented entry/exit criteria per stage, deals drift forward without real progress, making forecasts unreliable.",
        "mid": "Stage definitions exist but stage-to-stage conversion isn't reviewed on a regular cadence, so bottlenecks can go undiagnosed.",
        "high": "Enforced stage criteria with monthly conversion-rate review makes bottlenecks visible early, before they become a quarter-end surprise.",
    },
    "CON-OBJ-04": {
        "low": "Without a shared objection-handling playbook or battlecards, reps improvise inconsistently against common objections and competitors.",
        "mid": "Some enablement materials exist but aren't reinforced through coaching or call review.",
        "high": "A living playbook reinforced by coaching and call review keeps the whole team handling objections consistently, not just your best reps.",
    },
    "CON-WNL-05": {
        "low": "Without structured win/loss capture, the same avoidable losses tend to repeat because nobody is systematically learning from them.",
        "mid": "Win/loss reasons are captured but reviews aren't regular enough to turn findings into real process or coaching changes.",
        "high": "Regular win/loss reviews that feed into coaching and process changes turn every closed deal into a feedback loop that improves the next one.",
    },
    "CON-PGE-06": {
        "low": "Without hypothesis-driven testing on key conversion pages and forms, improvements are guesswork and regressions can go unnoticed.",
        "mid": "Some testing happens but isn't tracked with defined hypotheses or shared results, so learnings don't compound.",
        "high": "Structured, tracked experimentation on key conversion points compounds over time into a meaningfully better funnel.",
    },
    "DEL-TTV-01": {
        "low": "Without tracking time-to-first-value by segment, slow onboarding often goes unnoticed until it shows up as early churn.",
        "mid": "TTV is tracked in some form but there's no active initiative to shorten it, so it tends to drift rather than improve.",
        "high": "Actively tracked and shortened time-to-value is one of the strongest predictors of long-term retention.",
    },
    "DEL-ONB-02": {
        "low": "Without clear milestones, owners, and shared tracking, onboarding quality varies a lot between customers and between CSMs.",
        "mid": "A process exists but milestones or ownership aren't consistently clear across every account.",
        "high": "A tracked onboarding process with clear owners and realistic targets makes early customer experience consistent, not luck-dependent.",
    },
    "DEL-HLT-03": {
        "low": "Without a health-score model, at-risk accounts are usually only noticed after they've already started disengaging — too late for an effective save motion.",
        "mid": "Some health signals are tracked but there's no clear playbook for acting on them early.",
        "high": "A usage- and support-informed health score with a clear intervention playbook catches risk early enough to actually do something about it.",
    },
    "DEL-RET-04": {
        "low": "Low retention visibility usually means churn is discovered after the fact rather than predicted — teams at this stage typically don't have a defined at-risk signal yet.",
        "mid": "Retention is reviewed but not consistently by cohort or fast enough to launch targeted saves before a segment's decline compounds.",
        "high": "Monthly cohort-level retention review with fast, targeted response to decline is what turns retention from a lagging metric into a managed one.",
    },
    "DEL-QBR-05": {
        "low": "Without regular outcome-focused business reviews, high-value accounts can drift without anyone actively managing their trajectory or expansion potential.",
        "mid": "Reviews happen but aren't consistently tied to outcomes or roadmap alignment, limiting how much they actually influence the relationship.",
        "high": "Regular, outcome-focused QBRs linked to expansion pipeline keep your highest-value relationships actively managed, not just maintained.",
    },
    "DEL-ADV-06": {
        "low": "Without a structured way to capture testimonials and case studies, even happy customers rarely turn into usable proof points for future selling.",
        "mid": "Some advocacy content exists but isn't part of a consistent, ongoing programme.",
        "high": "A consistent advocacy programme turns customer success into a compounding sales asset — proof points that make the next deal easier.",
    },
}


def _score_guidance(id_code: str, score: int) -> str:
    """Look up the generic, pre-written guidance line for a question/score."""
    tiers = QUESTION_SCORE_GUIDANCE.get(id_code)
    if not tiers:
        return ""
    tier = "low" if score <= 2 else ("mid" if score == 3 else "high")
    return tiers.get(tier, "")


# ---------------------------------------------------------------------------
# Stage 1 — Text extractors
# ---------------------------------------------------------------------------

def _extract_csv(file_path: str) -> str:
    rows = []
    try:
        with open(file_path, newline='', encoding='utf-8', errors='replace') as f:
            reader = csv.reader(f)
            for i, row in enumerate(reader):
                if i > 500:
                    rows.append("... (truncated at 500 rows)")
                    break
                rows.append(", ".join(str(c) for c in row))
        return "\n".join(rows)
    except Exception as exc:
        logger.warning("CSV extraction failed for %s: %s", file_path, exc)
        return ""


def _extract_xlsx(file_path: str) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        parts = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            parts.append(f"[Sheet: {sheet_name}]")
            row_count = 0
            for row in ws.iter_rows(values_only=True):
                if row_count >= 300:
                    parts.append("... (truncated)")
                    break
                cells = [str(c) if c is not None else "" for c in row]
                if any(cells):
                    parts.append(", ".join(cells))
                    row_count += 1
        wb.close()
        return "\n".join(parts)
    except ImportError:
        logger.warning("openpyxl not installed; cannot parse XLSX %s", file_path)
        return ""
    except Exception as exc:
        logger.warning("XLSX extraction failed for %s: %s", file_path, exc)
        return ""


def _extract_pdf(file_path: str) -> str:
    try:
        from pypdf import PdfReader
        reader = PdfReader(file_path)
        pages = []
        for i, page in enumerate(reader.pages):
            if i >= 40:
                pages.append("... (truncated at 40 pages)")
                break
            text = page.extract_text() or ""
            if text:
                pages.append(text)
        return "\n".join(pages)
    except Exception as exc:
        logger.warning("PDF extraction failed for %s: %s", file_path, exc)
        return ""


def _extract_docx(file_path: str) -> str:
    try:
        import docx as _docx
        doc = _docx.Document(file_path)
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    except ImportError:
        pass
    except Exception as exc:
        logger.warning("python-docx failed for %s: %s", file_path, exc)

    try:
        with zipfile.ZipFile(file_path) as z:
            if 'word/document.xml' not in z.namelist():
                return ""
            with z.open('word/document.xml') as f:
                tree = ET.parse(f)
        ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        paragraphs = []
        for para in tree.findall('.//w:p', ns):
            texts = [node.text or '' for node in para.findall('.//w:t', ns)]
            text = ''.join(texts).strip()
            if text:
                paragraphs.append(text)
        return "\n".join(paragraphs)
    except Exception as exc:
        logger.warning("DOCX ZIP fallback failed for %s: %s", file_path, exc)
        return ""


def _extract_txt(file_path: str) -> str:
    try:
        with open(file_path, encoding='utf-8', errors='replace') as f:
            return f.read(_TEXT_LIMIT_PER_DOC)
    except Exception as exc:
        logger.warning("TXT extraction failed for %s: %s", file_path, exc)
        return ""


def _extract_json(file_path: str) -> str:
    try:
        with open(file_path, encoding='utf-8', errors='replace') as f:
            data = json.load(f)
        return json.dumps(data, indent=2)[:_TEXT_LIMIT_PER_DOC]
    except Exception as exc:
        logger.warning("JSON extraction failed for %s: %s", file_path, exc)
        return ""


def _extract_image_via_gemini(file_path: str, client, model_id: str, *, account=None) -> str:
    from .ai_services import _quota_cooldown_active, _request_budget_available, MONITORING_AVAILABLE
    if _quota_cooldown_active() or not _request_budget_available(account=account):
        logger.warning("AI unavailable (cooldown/credits) — skipping image extraction for %s", file_path)
        return ""
    try:
        from PIL import Image
        img = Image.open(file_path)
        response = client.models.generate_content(
            model=model_id,
            contents=[
                (
                    "Describe all content in this image in full detail — including every number, "
                    "metric, table value, chart label, and text string you can read. "
                    "Return plain text only."
                ),
                img,
            ]
        )
        if hasattr(response, "usage_metadata"):
            total_tokens = response.usage_metadata.total_token_count
            if MONITORING_AVAILABLE:
                from .utils_ai_monitoring import AIUsageTracker
                AIUsageTracker.log_usage(total_tokens, 'document_ocr')
            from .ai_credits import record_spend
            record_spend(account, total_tokens, "document_ocr")
        return (getattr(response, 'text', '') or '').strip()
    except Exception as exc:
        logger.warning("Gemini image extraction failed for %s: %s", file_path, exc)
        return ""


def extract_text_from_path(file_path: str, file_type: str,
                            client=None, model_id: str = "", *, account=None) -> str:
    """Route a saved file to the correct extractor and return extracted text."""
    if file_type == 'csv':
        text = _extract_csv(file_path)
    elif file_type == 'xlsx':
        text = _extract_xlsx(file_path)
    elif file_type == 'pdf':
        text = _extract_pdf(file_path)
    elif file_type == 'docx':
        text = _extract_docx(file_path)
    elif file_type == 'json':
        text = _extract_json(file_path)
    elif file_type == 'image':
        if client:
            text = _extract_image_via_gemini(file_path, client, model_id, account=account)
        else:
            text = ""
    else:  # txt, other
        text = _extract_txt(file_path)

    return text[:_TEXT_LIMIT_PER_DOC]


def extract_text_from_field(file_field, file_type: str,
                            client=None, model_id: str = "", *, account=None) -> str:
    """Backend-agnostic wrapper around extract_text_from_path.

    The format parsers need a real local filesystem path, but on cloud storage
    (GCS) FieldFile.path raises NotImplementedError("...doesn't support absolute
    paths"). Copy the stored file to a NamedTemporaryFile first so extraction
    works on any storage backend (local disk or GCS).
    """
    import os
    import tempfile

    suffix = os.path.splitext(file_field.name)[1]
    file_field.open('rb')
    try:
        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
            for chunk in file_field.chunks():
                tmp.write(chunk)
            tmp_path = tmp.name
    finally:
        file_field.close()

    try:
        return extract_text_from_path(tmp_path, file_type, client, model_id, account=account)
    finally:
        try:
            os.unlink(tmp_path)
        except OSError:
            pass


# ---------------------------------------------------------------------------
# Stage 2A — TF-IDF keyword ranking (scikit-learn)
# ---------------------------------------------------------------------------

def _tfidf_rank(chunks: list, queries: list, top_k: int = 3) -> dict:
    """
    Rank document paragraphs against each delivery question using TF-IDF cosine
    similarity.  Catches exact terminology: 'churn rate', 'onboarding milestone', etc.

    Returns {query_idx: [(chunk_text, score), ...]}
    """
    if not chunks or not queries:
        return {}
    try:
        from sklearn.feature_extraction.text import TfidfVectorizer
        from sklearn.metrics.pairwise import cosine_similarity
        import numpy as np

        corpus = chunks + queries
        vectorizer = TfidfVectorizer(
            stop_words="english",
            max_features=8000,
            ngram_range=(1, 2),
            sublinear_tf=True,
        )
        matrix     = vectorizer.fit_transform(corpus)
        chunk_vecs = matrix[: len(chunks)]
        query_vecs = matrix[len(chunks):]

        results = {}
        for qi in range(len(queries)):
            sims    = cosine_similarity(query_vecs[qi], chunk_vecs)[0]
            top_idx = np.argsort(sims)[::-1][:top_k]
            results[qi] = [
                (chunks[j], round(float(sims[j]), 4))
                for j in top_idx
                if sims[j] > 0.02
            ]
        return results

    except ImportError:
        logger.warning("scikit-learn not installed — TF-IDF stage skipped.")
        return {}
    except Exception as exc:
        logger.warning("TF-IDF ranking failed: %s", exc)
        return {}


# ---------------------------------------------------------------------------
# Stage 2B — Semantic similarity ranking (sentence-transformers)
# ---------------------------------------------------------------------------

def _get_sentence_transformer():
    """Lazy-load and cache the sentence-transformer model (all-MiniLM-L6-v2)."""
    global _ST_MODEL
    if _ST_MODEL is not None:
        return _ST_MODEL
    with _ST_LOCK:
        if _ST_MODEL is not None:
            return _ST_MODEL
        from sentence_transformers import SentenceTransformer
        logger.info(
            "Loading sentence-transformer model 'all-MiniLM-L6-v2' "
            "(first run downloads ~90 MB; subsequent runs use local cache)…"
        )
        _ST_MODEL = SentenceTransformer("all-MiniLM-L6-v2")
        logger.info("Sentence-transformer model ready.")
    return _ST_MODEL


def _semantic_rank(sentences: list, queries: list, top_k: int = 3) -> dict:
    """
    Rank document sentences against each delivery question using dense embeddings
    and cosine similarity.  Catches semantic paraphrases that TF-IDF misses —
    e.g. 'customers see results within a month' matches 'time to value' even
    without the exact phrase.

    Returns {query_idx: [(sentence_text, score), ...]}
    """
    if not sentences or not queries:
        return {}
    try:
        from sentence_transformers.util import cos_sim

        model      = _get_sentence_transformer()
        sent_embs  = model.encode(sentences, convert_to_tensor=True, show_progress_bar=False)
        query_embs = model.encode(queries,   convert_to_tensor=True, show_progress_bar=False)

        results = {}
        for qi in range(len(queries)):
            sims    = cos_sim(query_embs[qi], sent_embs)[0]
            top_idx = sims.argsort(descending=True)[:top_k]
            results[qi] = [
                (sentences[int(j)], round(float(sims[j]), 4))
                for j in top_idx
                if float(sims[j]) > 0.20
            ]
        return results

    except ImportError:
        logger.warning("sentence-transformers not installed — semantic ranking stage skipped.")
        return {}
    except Exception as exc:
        logger.warning("Semantic ranking failed: %s", exc)
        return {}


# ---------------------------------------------------------------------------
# Stage 2C — Named Entity Recognition (spaCy)
# ---------------------------------------------------------------------------

def _get_spacy_nlp():
    """Lazy-load and cache the spaCy en_core_web_sm pipeline.

    _extract_metrics_spacy (the only consumer of this cached pipeline) reads
    only doc.ents, never POS tags, the dependency parse, lemmas, or noun
    chunks -- so the tagger/parser/attribute_ruler/lemmatizer components are
    disabled, keeping just tok2vec + ner. This noticeably speeds up nlp()
    calls over the up-to-50k-char documents this pipeline processes.
    """
    global _SPACY_NLP
    if _SPACY_NLP is not None:
        return _SPACY_NLP
    with _SPACY_LOCK:
        if _SPACY_NLP is not None:
            return _SPACY_NLP
        try:
            import spacy
            _SPACY_NLP = spacy.load(
                "en_core_web_sm",
                disable=["tagger", "parser", "attribute_ruler", "lemmatizer"],
            )
        except ImportError:
            logger.warning("spaCy not installed — NER stage skipped.")
            return None
        except OSError:
            logger.warning(
                "spaCy model 'en_core_web_sm' not found. "
                "Run: python -m spacy download en_core_web_sm"
            )
            return None
    return _SPACY_NLP


def _extract_metrics_regex(text: str) -> list:
    """
    Regex-based metric extraction: catches percentages, money amounts, durations,
    dates, and common GTM KPI acronyms.  Used as primary output when spaCy is
    unavailable, and merged with spaCy results otherwise.
    """
    patterns = [
        (r'\b\d+(?:\.\d+)?%',                                    'PERCENT'),
        (r'\$[\d,]+(?:\.\d+)?(?:\s*[MBK](?:illion|illion)?)?',  'MONEY'),
        (r'\b\d+\s+(?:days?|weeks?|months?|years?)\b',           'DURATION'),
        (r'\bQ[1-4]\s*\d{4}\b',                                  'DATE'),
        (r'\b(?:FY|H[12])\s*\d{2,4}\b',                         'DATE'),
        (r'\b(?:NRR|GRR|CAC|LTV|ARR|MRR|TTV|NPS)\s*[:\-]?\s*'
         r'(?:\$?[\d,]+(?:\.\d+)?%?)',                           'KPI'),
        (r'\b\d{1,3}(?:,\d{3})*\s*(?:customers?|accounts?|users?|clients?)\b',
                                                                  'COUNT'),
    ]
    metrics, seen = [], set()
    for pattern, label in patterns:
        for m in re.finditer(pattern, text, re.IGNORECASE):
            key = m.group(0).lower().strip()
            if key in seen:
                continue
            seen.add(key)
            start = max(0, m.start() - 70)
            end   = min(len(text), m.end() + 70)
            ctx   = text[start:end].replace('\n', ' ').strip()
            metrics.append({'text': m.group(0), 'label': label, 'context': ctx})
    return metrics


def _extract_metrics_spacy(text: str) -> list:
    """
    Extract quantitative signals and key entities using spaCy NER, with automatic
    fallback to regex extraction when spaCy is unavailable or fails to load
    (e.g. numpy binary incompatibility — fix with: pip install spacy==3.8.3).

    Returns list of {"text", "label", "context"} dicts.
    """
    nlp = _get_spacy_nlp()
    if nlp is None:
        logger.info("spaCy unavailable — using regex metric extraction as fallback.")
        return _extract_metrics_regex(text)
    try:
        doc    = nlp(text[:50_000])
        wanted = {"PERCENT", "MONEY", "CARDINAL", "DATE", "TIME", "QUANTITY"}
        metrics, seen = [], set()

        for ent in doc.ents:
            if ent.label_ not in wanted:
                continue
            key = (ent.text.lower().strip(), ent.label_)
            if key in seen:
                continue
            seen.add(key)
            start = max(0, ent.start_char - 70)
            end   = min(len(text), ent.end_char + 70)
            ctx   = text[start:end].replace("\n", " ").strip()
            metrics.append({"text": ent.text, "label": ent.label_, "context": ctx})

        # Merge in regex results for KPI acronyms spaCy misses (NRR, GRR, TTV…)
        spacy_texts = {m['text'].lower() for m in metrics}
        for rm in _extract_metrics_regex(text):
            if rm['label'] == 'KPI' and rm['text'].lower() not in spacy_texts:
                metrics.append(rm)

        return metrics
    except Exception as exc:
        logger.warning("spaCy NER failed (%s) — falling back to regex extraction.", exc)
        return _extract_metrics_regex(text)


# ---------------------------------------------------------------------------
# Stage 2 orchestrator
# ---------------------------------------------------------------------------

def _clean_text_for_nlp(text: str) -> str:
    """Strip document separator headers before NLP processing."""
    return re.sub(r'=== Document: [^=]+ ===\n?', '', text)


def _split_paragraphs(text: str, min_len: int = 60) -> list:
    """Split text into meaningful paragraphs (double-newline boundaries)."""
    chunks = re.split(r'\n\s*\n', text)
    result = []
    for chunk in chunks:
        chunk = chunk.strip()
        if len(chunk) >= min_len:
            result.append(chunk)
        elif len(chunk) >= 20:
            for line in chunk.split('\n'):
                line = line.strip()
                if len(line) >= min_len:
                    result.append(line)
    return result


def _split_sentences(text: str, min_len: int = 25) -> list:
    """Lightweight sentence splitter (no NLTK required)."""
    sents = re.split(r'(?<=[.!?])\s+(?=[A-Z\"\'])', text)
    return [s.strip() for s in sents if len(s.strip()) >= min_len]


def _build_nlp_evidence(combined_text: str, questions: list = None) -> dict:
    """
    Orchestrate all three NLP/ML stages and return per-question structured evidence.

    When `questions` is None, defaults to DELIVERY_QUESTIONS.

    Schema:
        {
          "DEL-TTV-01": {
              "tfidf":    ["paragraph text …", …],
              "semantic": ["sentence text …",  …],
          },
          …,
          "_metrics": [{"text": "85%", "label": "PERCENT", "context": "…"}, …],
          "_stages_run": {"tfidf": bool, "semantic": bool, "ner": bool},
        }
    """
    if questions is None:
        questions = DELIVERY_QUESTIONS

    clean     = _clean_text_for_nlp(combined_text)
    paragraphs = _split_paragraphs(clean)
    sentences  = _split_sentences(clean)
    queries    = [f"{q['text']} {q['evidence_hint']}" for q in questions]

    logger.info(
        "NLP Stage 2 | %d paragraphs · %d sentences · %d questions",
        len(paragraphs), len(sentences), len(queries),
    )

    tfidf_results    = _tfidf_rank(paragraphs, queries, top_k=3)
    semantic_results = _semantic_rank(sentences, queries, top_k=3)
    ner_metrics      = _extract_metrics_spacy(clean)

    stages_run = {
        "tfidf":    bool(tfidf_results),
        "semantic": bool(semantic_results),
        "ner":      bool(ner_metrics),
    }
    logger.info(
        "NLP Stage 2 complete | TF-IDF: %s · Semantic: %s · NER: %d entities extracted",
        "yes" if stages_run["tfidf"]    else "no",
        "yes" if stages_run["semantic"] else "no",
        len(ner_metrics),
    )

    evidence = {"_metrics": ner_metrics, "_stages_run": stages_run}
    for qi, q in enumerate(questions):
        evidence[q["id_code"]] = {
            "tfidf":    [t for t, _ in tfidf_results.get(qi, [])],
            "semantic": [s for s, _ in semantic_results.get(qi, [])],
        }
    return evidence


# ---------------------------------------------------------------------------
# Stage 2.5 — Cross-company scoring memory (AI-unavailable fallback)
#
# Every successful Gemini scoring call captures the embedding of the passage
# it scored + the score it assigned into AIScoringExample. When Gemini is
# unavailable, _scoring_memory_fallback looks up the nearest past examples
# for the same question (any company) and uses a similarity-weighted vote as
# the suggested score — a real distillation of Gemini's own accumulated
# judgments, not a guess. Only embeddings + numeric scores are ever shared
# across companies; the evidence/reasoning shown to the user always comes
# from their own document, never another company's.
# ---------------------------------------------------------------------------

def _embed_one(text: str):
    """Embed a single passage with the shared sentence-transformer model."""
    model = _get_sentence_transformer()
    return model.encode([text], convert_to_numpy=True, show_progress_bar=False)[0]


def _capture_scoring_examples(nlp_evidence: dict, result: dict, questions: list):
    """Persist real Gemini scoring outcomes for reuse as a fallback signal.
    Never lets a capture failure affect the real analysis response."""
    from .models import AIScoringExample
    try:
        model = _get_sentence_transformer()
        rows = []
        for q in questions:
            q_id = q["id_code"]
            entry = result.get(q_id)
            if not entry or entry.get("score") is None:
                continue
            evidence = nlp_evidence.get(q_id, {})
            candidates = evidence.get("semantic") or evidence.get("tfidf") or []
            if not candidates:
                continue
            emb = model.encode([candidates[0]], convert_to_numpy=True, show_progress_bar=False)[0]
            rows.append(AIScoringExample(
                question_id_code=q_id,
                evidence_embedding=emb.tolist(),
                score=entry.get("score"),
                confidence=entry.get("confidence", ""),
            ))
        if rows:
            AIScoringExample.objects.bulk_create(rows)
    except Exception as exc:
        logger.warning("Failed to capture scoring examples: %s", exc)


def _nearest_scoring_examples(question_id_code: str, query_embedding, limit: int = 300, top_k: int = 5):
    """Return up to top_k (similarity, score) pairs for the given question,
    ranked by cosine similarity against the query embedding, drawn from the
    most recent `limit` examples in the shared corpus (any company)."""
    from .models import AIScoringExample
    import numpy as np

    rows = list(
        AIScoringExample.objects
        .filter(question_id_code=question_id_code)
        .order_by("-created_at")
        .values_list("evidence_embedding", "score")[:limit]
    )
    if not rows:
        return []

    embeddings = np.array([r[0] for r in rows])
    scores = [r[1] for r in rows]
    query = np.asarray(query_embedding)

    norms = np.linalg.norm(embeddings, axis=1) * (np.linalg.norm(query) or 1e-9)
    norms[norms == 0] = 1e-9
    sims = (embeddings @ query) / norms
    order = np.argsort(sims)[::-1][:top_k]
    return [(round(float(sims[i]), 4), scores[i]) for i in order if sims[i] >= 0.35]


def _scoring_memory_fallback(nlp_evidence: dict, questions: list) -> dict:
    """AI-unavailable substitute for a Gemini scoring call. For each question,
    finds this document's most relevant passage (already computed by
    _build_nlp_evidence) and looks up the nearest past Gemini scoring
    outcomes for that same question across the shared corpus, using a
    similarity-weighted vote as the suggested score.

    Falls back to a fixed, low-confidence score of 3 ("in progress") when the
    corpus has no examples yet for that question (cold start, before the
    corpus is seeded — see backfill_ai_scoring_examples) but this document
    does have *some* related evidence. Omits the question entirely when
    there's no supporting evidence at all — never invents a score with zero
    signal.
    """
    result = {}
    for q in questions:
        q_id = q["id_code"]
        evidence = nlp_evidence.get(q_id, {})
        candidates = evidence.get("semantic") or evidence.get("tfidf") or []
        passage = candidates[0] if candidates else None
        if not passage:
            continue  # no evidence at all — leave for manual scoring

        neighbors = []
        try:
            query_emb = _embed_one(passage)
            neighbors = _nearest_scoring_examples(q_id, query_emb)
        except Exception as exc:
            logger.warning("Scoring memory lookup failed for %s: %s", q_id, exc)

        if neighbors:
            weights = [sim for sim, _ in neighbors]
            weighted_score = sum(sim * score for sim, score in neighbors) / sum(weights)
            score = max(1, min(5, round(weighted_score)))
            score_spread = max(s for _, s in neighbors) - min(s for _, s in neighbors)
            confidence = "medium" if len(neighbors) >= 3 and score_spread <= 1 else "low"
        else:
            # Cold start for this question — corpus not seeded yet.
            score, confidence = 3, "low"

        excerpt = passage.strip()
        if len(excerpt) > 220:
            excerpt = excerpt[:220].rsplit(" ", 1)[0] + "…"

        result[q_id] = {
            "score": score,
            "confidence": confidence,
            "evidence": excerpt,
            "reasoning": (
                "AI scoring is temporarily unavailable — this suggestion is based on "
                "similar evidence patterns from past assessments. "
                + _score_guidance(q_id, score)
            ).strip(),
        }

    return result


# ---------------------------------------------------------------------------
# Stage 3 — Gemini analysis prompt (enriched with NLP evidence)
# ---------------------------------------------------------------------------

def _build_analysis_prompt(combined_text: str, nlp_evidence: dict = None,
                            questions: list = None, category_label: str = "Delivery") -> str:
    """
    Build the Gemini scoring prompt.  When nlp_evidence is provided, each
    question block includes the TF-IDF and semantic passages most relevant to
    that question, plus a global NER metrics section.  Falls back to raw text
    if NLP stages produced no output.
    """
    if questions is None:
        questions = DELIVERY_QUESTIONS
    question_parts = []
    for q in questions:
        id_code = q["id_code"]
        lines = [
            f"**{id_code} — {q['dimension']}**",
            f"Statement: \"{q['text']}\"",
            f"What to look for: {q['evidence_hint']}",
        ]

        if nlp_evidence and id_code in nlp_evidence:
            ev       = nlp_evidence[id_code]
            tfidf    = ev.get("tfidf", [])
            semantic = ev.get("semantic", [])
            if tfidf:
                lines.append("Keyword-matched passages (TF-IDF):")
                for p in tfidf[:2]:
                    lines.append(f"  • {p[:350]}")
            if semantic:
                lines.append("Semantically similar passages:")
                for s in semantic[:2]:
                    lines.append(f"  • {s[:350]}")
            if not tfidf and not semantic:
                lines.append("  (No strong evidence passages found for this question.)")

        question_parts.append("\n".join(lines))

    question_block = "\n\n".join(question_parts)

    # Global NER metrics section
    metrics_section = ""
    if nlp_evidence and nlp_evidence.get("_metrics"):
        key_labels  = {"PERCENT", "MONEY", "CARDINAL", "QUANTITY", "TIME"}
        key_metrics = [m for m in nlp_evidence["_metrics"] if m["label"] in key_labels][:25]
        if key_metrics:
            rows = [
                f"  [{m['label']}] {m['text']} — \"{m['context'][:100]}\""
                for m in key_metrics
            ]
            metrics_section = "\nEXTRACTED METRICS (spaCy NER):\n" + "\n".join(rows) + "\n"

    # NLP pipeline status note
    stages = (nlp_evidence or {}).get("_stages_run", {})
    active = [k for k, v in stages.items() if v]
    nlp_note = (
        f"\n[NLP pre-processing active: {', '.join(active)}]\n" if active else ""
    )

    # If NLP enrichment is working, a short raw excerpt suffices as fallback context;
    # otherwise include the full combined text so Gemini still has all information.
    if active:
        doc_section = (
            f"DOCUMENT EXCERPT (first 2 000 chars for reference):\n"
            f"---\n{combined_text[:2000]}\n---"
        )
    else:
        doc_section = f"UPLOADED DOCUMENTS:\n---\n{combined_text}\n---"

    n = len(questions)
    json_template = "\n".join(
        f'  "{q["id_code"]}": {{"score": <int 1-5>, "reasoning": "<one clear sentence>", '
        f'"evidence": "<exact quote or metric from document, or \'No direct evidence found\'>", '
        f'"confidence": "<high|medium|low>"}}'
        + ("," if i < n - 1 else "")
        for i, q in enumerate(questions)
    )

    return f"""You are a senior GTM analyst evaluating a company's {category_label} capabilities.

Your task: read the pre-processed evidence below and score the company on {n} {category_label} \
assessment statements using a 1–5 scale.
{nlp_note}
SCORING SCALE:
1 = No / Not in place
2 = Ad-hoc / Rarely
3 = In progress / Sometimes
4 = Consistent / Often
5 = Best-in-class / Always

SCORING GUIDANCE:
- Score ONLY based on evidence found in the documents. Do not assume.
- No evidence at all → score 1 or 2.
- Partial or informal evidence → score 2 or 3.
- Clear documented process but not measured/optimised → score 3 or 4.
- Measured, reviewed, and continuously improved → score 4 or 5.
- Be conservative: overconfident scores mislead the playbook generation.

{category_label.upper()} STATEMENTS WITH PRE-EXTRACTED EVIDENCE:
{question_block}
{metrics_section}
{doc_section}

Return ONLY a valid JSON object — no markdown, no extra text, no explanation outside the JSON:
{{
{json_template}
}}"""


# ---------------------------------------------------------------------------
# Public entry point
# ---------------------------------------------------------------------------

def analyze_delivery_documents(session) -> dict:
    """
    Three-stage pipeline:
      1. Collect extracted text from all DeliveryDocuments for this session.
      2. Run NLP/ML pre-processing (TF-IDF · sentence-transformers · spaCy NER).
      3. Call Gemini with the enriched, structured prompt.

    Returns a dict keyed by question id_code:
        {
          "DEL-TTV-01": {"score": 3, "reasoning": "…", "evidence": "…", "confidence": "medium"},
          …
        }
    Returns a dict with "_error" key on failure.
    """
    from .models import DeliveryDocument
    from .ai_services import (
        _get_client, _clean_json_response, _is_quota_error,
        _extract_retry_delay_seconds, _set_quota_cooldown, _quota_cooldown_active,
        _request_budget_available, MONITORING_AVAILABLE,
    )
    from .ai_credits import resolve_account_for_session, record_spend
    from .utils_ai_monitoring import AIUsageTracker
    from google.genai import types as genai_types

    docs = DeliveryDocument.objects.filter(session=session).exclude(extracted_text='')
    if not docs.exists():
        logger.warning("No documents with extracted text found for session %s", session.uuid)
        return {}

    combined_parts = []
    for doc in docs:
        combined_parts.append(
            f"=== Document: {doc.original_filename} ===\n{doc.extracted_text}"
        )
    combined_text = "\n\n".join(combined_parts)[:_COMBINED_TEXT_LIMIT]

    # ── Stage 2: NLP / ML pre-processing ────────────────────────────────────
    # Runs regardless of Gemini availability — it's also what the
    # AI-unavailable fallback (Stage 2.5) is built on.
    logger.info("Delivery analysis — starting NLP/ML pre-processing…")
    nlp_evidence = _build_nlp_evidence(combined_text, DELIVERY_QUESTIONS)

    account = resolve_account_for_session(session)

    client = _get_client()
    if not client or _quota_cooldown_active() or not _request_budget_available(account=account):
        logger.warning("Gemini unavailable for delivery analysis — using scoring memory fallback")
        return _scoring_memory_fallback(nlp_evidence, DELIVERY_QUESTIONS)

    # ── Stage 3: Gemini scoring with enriched prompt ─────────────────────────
    model_id = "gemini-2.5-flash"
    prompt   = _build_analysis_prompt(combined_text, nlp_evidence, DELIVERY_QUESTIONS, "Delivery")

    try:
        config = genai_types.GenerateContentConfig(
            temperature=0.2,
            max_output_tokens=2048,
            thinking_config=genai_types.ThinkingConfig(thinking_budget=0),
        )
        response = client.models.generate_content(
            model=model_id,
            contents=prompt,
            config=config,
        )
        if hasattr(response, "usage_metadata"):
            total_tokens = response.usage_metadata.total_token_count
            if MONITORING_AVAILABLE:
                AIUsageTracker.log_usage(total_tokens, 'delivery_analysis')
            record_spend(account, total_tokens, "delivery_analysis", session=session)

        raw     = getattr(response, 'text', '') or ''
        cleaned = _clean_json_response(raw)
        result  = json.loads(cleaned)

        expected_keys = {q['id_code'] for q in DELIVERY_QUESTIONS}
        if not expected_keys.issubset(result.keys()):
            logger.warning(
                "Gemini delivery analysis returned incomplete keys: %s", list(result.keys())
            )
            return {}

        for key in expected_keys:
            entry = result[key]
            entry['score'] = max(1, min(5, int(entry.get('score', 1))))
            if entry.get('confidence') not in ('high', 'medium', 'low'):
                entry['confidence'] = 'medium'

        _capture_scoring_examples(nlp_evidence, result, DELIVERY_QUESTIONS)
        return result

    except json.JSONDecodeError as exc:
        logger.warning("Delivery analysis JSON parse error: %s | raw[:300]=%s", exc, raw[:300])
        return {"_error": f"JSON parse failed: {exc}"}
    except Exception as exc:
        logger.warning("Delivery analysis Gemini error: %s", exc)
        if _is_quota_error(exc):
            _set_quota_cooldown(_extract_retry_delay_seconds(exc))
        fallback = _scoring_memory_fallback(nlp_evidence, DELIVERY_QUESTIONS)
        if fallback:
            return fallback
        return {"_error": str(exc)}
