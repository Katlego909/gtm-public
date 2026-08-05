# gtm/utils_docx.py
import re
from io import BytesIO
from django.contrib.staticfiles import finders
from django.http import HttpResponse
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

from .utils_pdf import _normalize_ai_markdown, _parse_playbook_priorities, _slugify_filename_part

# Same burnt-copper brand accent as the PDF export (theme/static/css/app.css's
# --accent-600 / --accent-900) -- kept as RGBColor tuples since python-docx
# doesn't take hex strings directly.
_ACCENT_RGB = RGBColor(0xB8, 0x53, 0x0F)
_ACCENT_DARK_RGB = RGBColor(0x5C, 0x2B, 0x10)

_LOGO_PATH = finders.find("images/forge_logo.png")


def render_insight_docx_response(*, company_name, ai_playbook_md, doc_title=None, doc_type_label=None, doc_date=None):
    """
    Build a Word (.docx) export of a single AI insight, structurally matching
    the PDF export: a title, then each parsed "Priority N: Title (Category)"
    section as a heading with its bullet points.

    doc_title/doc_type_label/doc_date: see render_insight_pdf_response's
    docstring in gtm/utils_pdf.py -- same optional-identity filename scheme.

    Note: unlike the PDF (which embeds its heading font directly, so it
    always renders correctly), a .docx only stores a font *name* -- Word
    would render a custom typeface only if the recipient's machine has it
    installed, which can't be guaranteed. Headings stay on Word's default
    font; only brand color and the logo are reliably portable here.
    """
    normalized = _normalize_ai_markdown(ai_playbook_md or "")
    priorities = _parse_playbook_priorities(normalized)

    doc = Document()

    if _LOGO_PATH:
        logo_paragraph = doc.add_paragraph()
        logo_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        logo_run = logo_paragraph.add_run()
        try:
            logo_run.add_picture(_LOGO_PATH, width=Cm(4.5))
        except Exception:
            pass

    title = doc.add_heading("Funti3r GTM Validator", level=0)
    for run in title.runs:
        run.font.color.rgb = _ACCENT_DARK_RGB

    subtitle = doc.add_heading(f"{company_name or 'Company'} — AI Insight", level=1)
    for run in subtitle.runs:
        run.font.color.rgb = _ACCENT_RGB

    if priorities:
        for priority in priorities:
            heading_text = priority["title"]
            if priority.get("category"):
                heading_text = f"{heading_text} ({priority['category']})"
            doc.add_heading(heading_text, level=2)
            for bullet in priority.get("bullets", []):
                doc.add_paragraph(bullet, style="List Bullet")
    else:
        # Fallback: no structured "Priority N" sections found — dump the
        # normalized markdown as plain paragraphs so the export is never empty.
        plain_text = re.sub(r"[*_#>`-]", "", normalized).strip()
        for line in plain_text.splitlines():
            line = line.strip()
            if line:
                doc.add_paragraph(line)

    buffer = BytesIO()
    doc.save(buffer)
    docx_bytes = buffer.getvalue()
    buffer.close()

    resp = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    company_slug = _slugify_filename_part(company_name, fallback="company")
    _parts = [company_slug]
    if doc_title:
        _parts.append(_slugify_filename_part(doc_title, fallback="document", max_len=50))
    _parts.append(_slugify_filename_part(doc_type_label, fallback="insight") if doc_type_label else "insight")
    if doc_date:
        _parts.append(doc_date.strftime("%Y%m%d"))
    _parts.append("ForgeGTM")
    _filename = "_".join(_parts) + ".docx"
    resp["Content-Disposition"] = f'attachment; filename="{_filename}"'
    resp.write(docx_bytes)
    return resp
